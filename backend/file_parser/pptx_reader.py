"""
PPT 解析器 — 独立封装

仅依赖 python-pptx。
提取所有形状的文本+表格（不限于 title 占位符，适配中文 PPT 布局）。
"""

import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation


def read_pptx(filepath: str | Path) -> list[dict[str, Any]]:
    prs = Presentation(str(filepath))
    slides: list[dict[str, Any]] = []

    for i, slide in enumerate(prs.slides, 1):
        item: dict[str, Any] = {
            "slide_number": i,
            "texts": [],
            "tables": [],
            "notes": "",
        }

        # 提取所有形状的文本（不限于 title 占位符）
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        item["texts"].append(text)

            if shape.has_table:
                table_rows = []
                for row in shape.table.rows:
                    table_rows.append([cell.text.strip() for cell in row.cells])
                if table_rows:
                    item["tables"].append(table_rows)

        # 首段作为标题
        if item["texts"]:
            item["title"] = item["texts"][0]

        # 备注
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                item["notes"] = notes_text

        slides.append(item)

    return slides


def main():
    if len(sys.argv) < 2:
        print("用法: python -m file_parser.pptx_reader <PPTX文件路径>")
        sys.exit(1)

    path = sys.argv[1]
    if not Path(path).exists():
        print(f"文件不存在: {path}")
        sys.exit(1)

    data = read_pptx(path)
    for s in data:
        print(f"Slide {s['slide_number']}: {s.get('title', '')[:60]}")
        if s["texts"]:
            print(f"  文本: {len(s['texts'])} 段")
        if s["tables"]:
            print(f"  表格: {len(s['tables'])} 个")


if __name__ == "__main__":
    main()
