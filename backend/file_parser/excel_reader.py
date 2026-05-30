"""
Excel 解析器 — 独立封装

仅依赖 openpyxl，可脱离 Web 框架运行。
use read_only=True 流式解析，速度比默认模式快 3-5x。
"""

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl


def read_excel(filepath: str | Path, sheets: list[str] | None = None) -> dict[str, list[dict[str, Any]]]:
    """
    流式读取 Excel，返回 { sheet_name: [{col: val}, ...] }
    - read_only=True: 流式 XML 解析，不加载完整对象模型
    - sheets: 指定要读取的 sheet 名列表，None=全部
    """
    wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
    result: dict[str, list[dict[str, Any]]] = {}
    target = sheets if sheets else wb.sheetnames

    for name in target:
        if name not in wb.sheetnames:
            result[name] = []
            continue

        ws = wb[name]
        rows_iter = ws.iter_rows(values_only=True)

        try:
            headers_row = next(rows_iter)
        except StopIteration:
            result[name] = []
            continue

        headers = [str(h).strip() if h else f"col_{i}" for i, h in enumerate(headers_row)]
        ncols = len(headers)
        data: list[dict[str, Any]] = []

        for row in rows_iter:
            if not row or all(v is None for v in row):
                continue
            record = {}
            for i in range(min(len(row), ncols)):
                if row[i] is not None:
                    record[headers[i]] = row[i]
            if record:
                data.append(record)

        result[name] = data

    wb.close()
    return result


def main():
    if len(sys.argv) < 2:
        print("用法: python -m file_parser.excel_reader <Excel文件路径> [sheet1,sheet2,...]")
        sys.exit(1)

    path = sys.argv[1]
    target = sys.argv[2].split(",") if len(sys.argv) > 2 else None

    if not Path(path).exists():
        print(f"文件不存在: {path}")
        sys.exit(1)

    data = read_excel(path, target)
    print(json.dumps({k: f"{len(v)} 行" for k, v in data.items()}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
